"""
Offline tests for the Week 4 submission workflow: evidence, screening rules, requirement parsing,
drafting, the grounding check, approval and evaluation. No Azure calls.

Run: pytest -q
"""
import asyncio
import json
import sys
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(Path(__file__).parent))

from copilot.evidence import evidence_for, mentions  # noqa: E402
from copilot.profiles import fix_ai_misreads, to_profile  # noqa: E402
from copilot.requirements import (compare_requirement, load_requirements, normalize_skill,  # noqa: E402
                                  skill_vocabulary, to_requirement)
from copilot.screening import evaluate_criterion, screen  # noqa: E402
from copilot.submission import (ApprovalRequired, build_draft, check_grounded, export_docx, review,  # noqa: E402
                                save_draft, to_markdown, write_summary, _facts)
from copilot.submission_eval import compare_screening  # noqa: E402
from test_extraction_pipeline import GT_PATH, arr, fake_cu_result, obj, s  # noqa: E402

TODAY = date(2026, 10, 1)


# ---------------------------------------------------------------------------
# Fixtures: what Content Understanding would return if it read everything perfectly
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def gt():
    if not GT_PATH.exists() or "requests" not in json.loads(GT_PATH.read_text()):
        pytest.skip("Run scripts/generate_synthetic_data.py first")
    return json.loads(GT_PATH.read_text())


def cu_with_highlights(c: dict) -> dict:
    result = fake_cu_result(c)
    result["contents"][0]["fields"]["WorkHistory"] = arr([obj({
        "JobTitle": s(j["JobTitle"]), "Company": s(j["Company"]), "StartDate": s(j["StartDate"]),
        "EndDate": s(j["EndDate"]), "Highlights": arr([s(h) for h in j["Highlights"]])}) for j in c["experience"]])
    return result


@pytest.fixture(scope="module")
def profiles(gt):
    return {c["id"]: to_profile(c["id"], c["file"], cu_with_highlights(c)) for c in gt["candidates"]}


def cu_requirement(r: dict) -> dict:
    """A perfect requirement-analyzer result for one RFR."""
    num = lambda v: {"type": "number", "valueNumber": v}  # noqa: E731
    fields = {
        "RequestId": s(r["id"]), "Title": s(r["title"]), "Client": s(r["client"]), "Term": s(r["term"]),
        "Location": s(r["location"]), "Positions": {"type": "integer", "valueInteger": r["positions"]},
        "MandatoryCriteria": arr([obj({"CriterionId": s(c["id"]), "Text": s(c["text"]),
                                       "Skills": arr([s(x) for x in c["skills"]]),
                                       "MinimumYears": num(c["min_years"]) if c["min_years"] else None})
                                  for c in r["mandatory"]]),
        "RatedCriteria": arr([obj({"CriterionId": s(c["id"]), "Text": s(c["text"]),
                                   "Skills": arr([s(x) for x in c["skills"]]),
                                   "TargetYears": num(c["target_years"]), "Points": num(c["points"])})
                              for c in r["rated"]]),
    }
    return {"contents": [{"kind": "document", "fields": fields}]}


# ---------------------------------------------------------------------------
# Evidence
# ---------------------------------------------------------------------------
def test_mentions_whole_terms():
    assert mentions("Built CI/CD pipelines in GitHub Actions", "CI/CD")
    assert mentions("Services in C# and .NET", "C#") and mentions("Services in C# and .NET", ".NET")
    assert not mentions("Worked at Google", "Go")
    assert not mentions("Used Kafka Streams", "")


def test_evidence_counts_overlap_once():
    jobs = [{"JobTitle": "A", "Company": "X", "StartDate": "Jan 2020", "EndDate": "Dec 2022",
             "Highlights": ["Built pipelines in Apache Spark"]},
            {"JobTitle": "B", "Company": "Y", "StartDate": "Jan 2022", "EndDate": "Dec 2023",
             "Highlights": ["Tuned Spark jobs", "Ran Apache Spark on Databricks"]},
            {"JobTitle": "C", "Company": "Z", "StartDate": "sometime", "EndDate": "Present",
             "Highlights": ["Apache Spark again"]}]
    ev = evidence_for(jobs, ["Apache Spark"], TODAY)
    assert [j["company"] for j in ev["jobs"]] == ["X", "Y", "Z"]
    assert ev["months"] == 47                         # Jan 2020 - Dec 2023, overlap once; Z has no dates
    assert ev["jobs"][1]["highlights"] == ["Ran Apache Spark on Databricks"]


# ---------------------------------------------------------------------------
# Screening rules
# ---------------------------------------------------------------------------
PROFILE = {"id": "C1", "skills": ["Terraform", "Docker"], "work_history": [
    {"JobTitle": "Cloud Engineer", "Company": "Acme", "StartDate": "Oct 2023", "EndDate": "Present",
     "Highlights": ["Provisioned 20+ environments with Terraform"]}]}  # 3 years


def test_mandatory_met_not_met_verify():
    ok = evaluate_criterion({"id": "M1", "skills": ["Terraform"], "min_years": 3}, "mandatory", PROFILE, TODAY)
    short = evaluate_criterion({"id": "M1", "skills": ["Terraform"], "min_years": 4}, "mandatory", PROFILE, TODAY)
    listed = evaluate_criterion({"id": "M2", "skills": ["Docker"], "min_years": 0}, "mandatory", PROFILE, TODAY)
    none = evaluate_criterion({"id": "M3", "skills": ["Kubernetes"], "min_years": 0}, "mandatory", PROFILE, TODAY)
    assert (ok["status"], ok["evidence_years"]) == ("met", 3.0)
    assert short["status"] == "not_met"
    assert listed["status"] == "verify" and listed["listed_in_skills"] == ["Docker"]
    assert none["status"] == "not_met"


def test_rated_points_are_proportional():
    half = evaluate_criterion({"id": "R1", "skills": ["Terraform"], "target_years": 6, "points": 40}, "rated", PROFILE, TODAY)
    full = evaluate_criterion({"id": "R1", "skills": ["Terraform"], "target_years": 2, "points": 40}, "rated", PROFILE, TODAY)
    listed = evaluate_criterion({"id": "R2", "skills": ["Docker"], "target_years": 2, "points": 20}, "rated", PROFILE, TODAY)
    assert half["points_awarded"] == 20 and half["status"] == "partial"
    assert full["points_awarded"] == 40 and full["status"] == "met"
    assert listed["points_awarded"] == 0 and listed["status"] == "verify"


def test_mandatory_beats_points(gt, profiles):
    """A high scorer who fails a mandatory must rank below every compliant candidate."""
    req = load_requirements("truth")["RFR-003"]   # CAND-019 has full points but only 2.4 of 3 years for M1
    res = screen(req, list(profiles.values()), TODAY)
    order = [r["compliance"] for r in res]
    assert order == sorted(order, key={"met": 0, "verify": 1, "not_met": 2}.get)
    failed_high = max((r for r in res if r["compliance"] == "not_met"), key=lambda r: r["points"])
    met = [r for r in res if r["compliance"] == "met"]
    assert failed_high["points"] >= min(r["points"] for r in met)   # as many points as a compliant candidate...
    assert res.index(failed_high) > max(res.index(r) for r in met)   # ...but still ranked below all of them


def test_perfect_extraction_reproduces_ground_truth(gt, profiles):
    reqs = load_requirements("truth")
    for r in gt["requests"]:
        m = compare_screening(screen(reqs[r["id"]], list(profiles.values()), TODAY), r["expected"])
        assert m["compliance_accuracy"] == 1 and m["points_within_5"] == 1, r["id"]
        assert not m["false_pass"] and not m["false_fail"] and m["shortlist_overlap"] == 1


def test_eval_catches_a_false_pass(gt, profiles):
    r = gt["requests"][0]
    res = screen(load_requirements("truth")[r["id"]], list(profiles.values()), TODAY)
    wrong = next(x for x in res if x["compliance"] == "not_met")
    wrong["compliance"] = "met"
    m = compare_screening(res, r["expected"])
    assert m["false_pass"] == [wrong["candidate_id"]] and m["compliance_accuracy"] < 1


# ---------------------------------------------------------------------------
# Requirement extraction
# ---------------------------------------------------------------------------
def test_requirement_parsing_round_trip(gt, profiles):
    vocab = skill_vocabulary(profiles)
    for r in gt["requests"]:
        parsed = to_requirement(r["id"], cu_requirement(r), vocab)
        assert compare_requirement(parsed, r)["accuracy"] == 1, r["id"]
        assert parsed["title"] == r["title"] and parsed["positions"] == r["positions"]
    assert to_requirement("RFR-004", cu_requirement(gt["requests"][3]), vocab)["mandatory"][0]["min_years"] == 2


def test_skill_normalization():
    vocab = ["Azure Data Factory", "CI/CD", "Kafka"]
    assert normalize_skill("azure datafactory", vocab) == "Azure Data Factory"
    assert normalize_skill("CI/CD", vocab) == "CI/CD"
    assert normalize_skill("COBOL", vocab) == "COBOL"          # unknown terms are kept


def test_requirement_comparison_flags_errors(gt):
    r = gt["requests"][0]
    bad = json.loads(json.dumps({k: v for k, v in r.items() if k != "expected"}))
    bad["mandatory"][0]["min_years"] = 5
    bad["rated"].pop()
    res = compare_requirement(bad, r)
    assert not res["checks"]["M1 min_years"] and not res["checks"]["R3 found"] and res["accuracy"] < 1


# ---------------------------------------------------------------------------
# Drafts
# ---------------------------------------------------------------------------
@pytest.fixture(scope="module")
def drafts(gt, profiles):
    req = load_requirements("truth")["RFR-002"]
    res = screen(req, list(profiles.values()), TODAY)
    met = next(r for r in res if r["compliance"] == "met")
    ver = next(r for r in res if r["compliance"] == "verify")
    return (build_draft(req, met, profiles[met["candidate_id"]]),
            build_draft(req, ver, profiles[ver["candidate_id"]]))


def test_draft_cites_evidence_and_is_blind(drafts, profiles):
    good, verify = drafts
    m1 = good["rows"][0]
    assert m1["status"] == "met" and m1["evidence"][0]["company"] in m1["response"]
    assert m1["evidence"][0]["start"] in m1["response"]
    assert not good["warnings"]
    assert verify["warnings"] and "listed but not shown" in verify["warnings"][0]
    assert "needs verification" in verify["summary"]
    for d in drafts:
        p = profiles[d["candidate_id"]]
        text = json.dumps(d)
        assert p["name"] not in text and p["email"] not in text
    assert "| M1 | met |" in to_markdown(good)


def test_grounding_check(drafts, profiles):
    facts = _facts(drafts[0])
    vocab = skill_vocabulary(profiles)
    yrs = drafts[0]["rows"][0]["evidence_years"]
    assert check_grounded(f"The candidate shows {yrs:g} years of relevant experience and meets all mandatory "
                          "requirements.", facts, vocab) == []
    assert "number not in evidence: 97" in check_grounded("Scores 97 points.", facts, vocab)
    unseen = next(v for v in vocab if v.lower() not in json.dumps(facts).lower())
    assert f"skill not in evidence: {unseen}" in check_grounded(f"Strong in {unseen}.", facts, vocab)
    assert check_grounded("Young and single, a great fit.", facts, vocab)


def _scripted(text):
    from agent_framework import BaseChatClient, ChatResponse, Message

    class Scripted(BaseChatClient):
        async def _inner_get_response(self, *, messages, stream=False, options=None, **kwargs):
            return ChatResponse(messages=[Message(role="assistant", contents=[text])])
    return Scripted()


def test_model_summary_kept_only_when_grounded(drafts, profiles):
    pytest.importorskip("agent_framework")
    vocab = skill_vocabulary(profiles)
    good = json.loads(json.dumps(drafts[0]))
    asyncio.run(write_summary(good, _scripted("The candidate meets every mandatory requirement."), vocab))
    assert good["summary_source"] == "model"
    bad = json.loads(json.dumps(drafts[0]))
    template = bad["summary"]
    asyncio.run(write_summary(bad, _scripted("The candidate has 25 years with Snowflake."), vocab))
    assert bad["summary_source"] == "template" and bad["summary"] == template
    assert bad["summary_rejected"]["problems"]


def test_export_requires_approval(drafts, profiles, tmp_path):
    draft = drafts[0]
    profile = profiles[draft["candidate_id"]]
    path = save_draft(json.loads(json.dumps(draft)), tmp_path)
    assert path.with_suffix(".md").exists()
    with pytest.raises(ApprovalRequired):
        export_docx(json.loads(path.read_text()), profile, tmp_path / "out.docx")
    with pytest.raises(ValueError):
        review(path, "  ", "approved")
    with pytest.raises(ValueError):
        review(path, "Yusuf", "maybe")
    rejected = review(path, "Yusuf", "rejected", "not a fit")
    with pytest.raises(ApprovalRequired):
        export_docx(rejected, profile, tmp_path / "out.docx")
    approved = review(path, "Yusuf", "approved")
    out = export_docx(approved, profile, tmp_path / "out.docx")
    from docx import Document
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert profile["name"] in text and "approved by Yusuf" in text   # name appears only after approval
    assert len(doc.tables[0].rows) == 1 + len(draft["rows"])


# ---------------------------------------------------------------------------
# Agent tools
# ---------------------------------------------------------------------------
def test_screen_request_tool(profiles):
    from copilot.agent_tools import RecruitingTools, keyword_search_backend

    tools = RecruitingTools(profiles, {}, keyword_search_backend(profiles), load_requirements("truth"))
    assert {r["request_id"] for r in tools.list_requests()} == {"RFR-001", "RFR-002", "RFR-003", "RFR-004"}
    out = tools.screen_request("rfr-003", top=3)
    assert len(out["candidates"]) == 3 and out["candidates"][0]["compliance"] == "met"
    assert sum(out["counts"].values()) == len(profiles)
    assert not any(p["name"] in json.dumps(out) for p in profiles.values())
    assert "error" in tools.screen_request("RFR-999")


# ---------------------------------------------------------------------------
# Fixes from the first live run (Oct 21, 2026)
# ---------------------------------------------------------------------------
def test_ai_misreads_fixed_in_highlights():
    # real Content Understanding output for CAND-010 and CAND-028
    assert fix_ai_misreads("Deployed a RAG assistant on Azure Al Search serving 44K queries per month") == \
        "Deployed a RAG assistant on Azure AI Search serving 44K queries per month"
    assert fix_ai_misreads("Deployed a RAG assistant on Azure OpenAl serving 6K queries") == \
        "Deployed a RAG assistant on Azure OpenAI serving 6K queries"
    assert fix_ai_misreads("Built agents in Al Foundry") == "Built agents in AI Foundry"
    assert fix_ai_misreads("Pair-programmed with Al on the billing API") == "Pair-programmed with Al on the billing API"
    raw = {"contents": [{"fields": {"WorkHistory": arr([obj({
        "JobTitle": s("ML Engineer"), "Company": s("X"), "StartDate": s("Jan 2020"), "EndDate": s("Jan 2023"),
        "Highlights": arr([s("Shipped search on Azure Al Search")])})])}}]}
    prof = to_profile("C9", "f", raw)
    assert evidence_for(prof["work_history"], ["Azure AI Search"], TODAY)["months"] == 36


def test_skill_normalization_contained_term():
    assert normalize_skill("Zero Trust policies", ["Zero Trust", "Entra ID"]) == "Zero Trust"
    assert normalize_skill("Azure AI Search index", ["Azure", "Azure AI Search"]) == "Azure AI Search"  # longest


def test_grounding_allows_fair_rounding(drafts, profiles):
    """The live model wrote 'more than 5 years' for 5.2 years of evidence; that is true and must pass."""
    facts = _facts(drafts[0])
    vocab = skill_vocabulary(profiles)
    yrs = max(r["evidence_years"] for r in drafts[0]["rows"])
    assert check_grounded(f"The candidate has more than {int(yrs)} years of relevant experience.", facts, vocab) == []
    assert check_grounded(f"The candidate has {int(yrs) + 3} years of relevant experience.", facts, vocab)


def test_requirement_title_is_checked(gt):
    r = gt["requests"][1]
    got = {k: v for k, v in r.items() if k != "expected"} | {"title": "Developer"}  # what the live run returned
    assert not compare_requirement(got, r)["checks"]["title"]


def test_stale_profiles_detected(profiles):
    from copilot.agent_tools import stale_profiles

    assert stale_profiles(profiles) == []
    old = {"CAND-019": {"work_history": [{"JobTitle": "Cloud Engineer", "StartDate": "Jan 2020", "EndDate": "Present"}]}}
    assert stale_profiles(profiles | old) == ["CAND-019"]

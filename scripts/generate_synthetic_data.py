"""
Generate a synthetic recruiting dataset: resumes (PDF / DOCX / scanned-style PNG),
job descriptions, public-sector style Requests for Resumes, and ground-truth labels for evaluation.

All people and companies are fake (Faker). Never put real candidate data in this repo.

Usage:
    python scripts/generate_synthetic_data.py              # 40 resumes, seed 42
    python scripts/generate_synthetic_data.py --count 60 --seed 7
"""
from __future__ import annotations

import argparse
import json
import random
import shutil
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt
from faker import Faker
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "synthetic"
TODAY = date(2026, 10, 1)

# Mixed locales -> a realistic spread of names, which lets us test blind screening later.
LOCALES = ["en_CA", "en_IN", "fr_CA", "en_GB", "es_MX", "de_DE", "it_IT", "en_US", "pt_BR", "nl_NL"]
CITIES = ["Toronto, ON", "Mississauga, ON", "Ottawa, ON", "Waterloo, ON", "Montreal, QC",
          "Vancouver, BC", "Calgary, AB", "Halifax, NS", "Brampton, ON", "Markham, ON"]

# ---------------------------------------------------------------------------
# Role families
# ---------------------------------------------------------------------------
FAMILIES = {
    "data_engineer": {
        "titles": ["Data Analyst", "Data Engineer", "Senior Data Engineer", "Lead Data Engineer"],
        "core": ["Python", "SQL", "Apache Spark", "Azure Data Factory", "Databricks", "Airflow"],
        "extra": ["Kafka", "dbt", "Snowflake", "Delta Lake", "Synapse Analytics", "Power BI",
                  "Terraform", "Docker", "Scala", "Microsoft Fabric"],
        "certs": ["Microsoft Certified: Fabric Data Engineer Associate",
                  "Databricks Certified Data Engineer Associate"],
        "bullets": [
            ("Built batch and streaming pipelines in {a} processing {n}M+ records per day", ["Apache Spark", "Databricks", "Kafka", "Python"]),
            ("Migrated legacy SSIS jobs to {a}, cutting run time by {p}%", ["Azure Data Factory", "Airflow", "Microsoft Fabric"]),
            ("Designed a medallion lakehouse on {a} used by {n} analytics teams", ["Databricks", "Delta Lake", "Microsoft Fabric", "Snowflake"]),
            ("Implemented data quality checks with {a}, reducing incidents by {p}%", ["dbt", "Python", "SQL"]),
        ],
    },
    "ml_engineer": {
        "titles": ["Junior ML Engineer", "Machine Learning Engineer", "Senior ML Engineer", "AI Engineer"],
        "core": ["Python", "PyTorch", "Azure Machine Learning", "MLOps", "Docker", "LLMs", "Azure OpenAI", "RAG"],
        "extra": ["LangChain", "Azure AI Search", "scikit-learn", "Kubernetes",
                  "MLflow", "Prompt Engineering", "Microsoft Foundry", "Hugging Face", "FastAPI"],
        "certs": ["Microsoft Certified: Azure AI Engineer Associate",
                  "Microsoft Certified: Azure Data Scientist Associate"],
        "bullets": [
            ("Deployed a RAG assistant on {a} serving {n}K queries per month", ["Azure OpenAI", "Azure AI Search", "LangChain", "Microsoft Foundry"]),
            ("Reduced model inference latency by {p}% by containerizing with {a}", ["Docker", "Kubernetes", "FastAPI"]),
            ("Built an evaluation harness for LLM outputs, raising groundedness by {p}%", []),
            ("Automated model retraining pipelines with {a} across {n} production models", ["Azure Machine Learning", "MLflow", "MLOps"]),
            ("Trained and fine-tuned deep learning models in {a}, improving accuracy by {p}%", ["PyTorch", "Hugging Face", "scikit-learn"]),
        ],
    },
    "cloud_devops": {
        "titles": ["Systems Administrator", "Cloud Engineer", "DevOps Engineer", "Senior Platform Engineer"],
        "core": ["Azure", "Terraform", "Kubernetes", "CI/CD", "Linux", "Docker"],
        "extra": ["AKS", "GitHub Actions", "Azure DevOps", "Bicep", "Prometheus", "Grafana", "Helm",
                  "PowerShell", "Bash", "AWS", "Ansible"],
        "certs": ["Microsoft Certified: Azure Administrator Associate",
                  "Microsoft Certified: DevOps Engineer Expert",
                  "HashiCorp Certified: Terraform Associate", "Certified Kubernetes Administrator"],
        "bullets": [
            ("Provisioned {n}+ cloud environments with {a}, eliminating manual setup", ["Terraform", "Bicep", "Ansible"]),
            ("Cut deployment time by {p}% by rebuilding pipelines in {a}", ["GitHub Actions", "Azure DevOps", "CI/CD"]),
            ("Ran {a} clusters hosting {n} microservices at 99.9% uptime", ["AKS", "Kubernetes"]),
            ("Introduced monitoring dashboards in {a}, cutting incident detection time by {p}%", ["Grafana", "Prometheus"]),
        ],
    },
    "backend": {
        "titles": ["Junior Developer", "Software Developer", "Senior Software Engineer", "Staff Engineer"],
        "core": ["Java", "Spring Boot", "REST APIs", "SQL", "Microservices", "Git"],
        "extra": ["C#", ".NET", "Python", "Node.js", "PostgreSQL", "Redis", "Kafka", "Azure",
                  "Docker", "gRPC", "Go"],
        "certs": ["Oracle Certified Professional: Java SE Developer",
                  "Microsoft Certified: Azure Developer Associate"],
        "bullets": [
            ("Designed {a} services handling {n}K requests per minute", ["Spring Boot", ".NET", "Node.js", "Go"]),
            ("Refactored a monolith into microservices, reducing release cycle by {p}%", []),
            ("Improved API response time by {p}% using {a} caching", ["Redis"]),
            ("Mentored {n} junior developers on {a} best practices", ["Java", "C#", "Python", "Git"]),
        ],
    },
    "frontend": {
        "titles": ["Web Developer", "Frontend Developer", "Senior Frontend Engineer", "UI Engineering Lead"],
        "core": ["JavaScript", "TypeScript", "React", "HTML", "CSS", "REST APIs"],
        "extra": ["Next.js", "Redux", "Tailwind CSS", "Jest", "Playwright", "Figma", "GraphQL",
                  "Vue.js", "Accessibility (WCAG)", "Storybook"],
        "certs": ["Meta Front-End Developer Professional Certificate"],
        "bullets": [
            ("Rebuilt the customer portal in {a}, improving page load by {p}%", ["React", "Next.js", "Vue.js"]),
            ("Built a shared component library in {a} used by {n} product teams", ["Storybook", "React", "TypeScript"]),
            ("Raised test coverage to {p}% with {a}", ["Jest", "Playwright"]),
            ("Led accessibility (WCAG) remediation across {n} pages", []),
        ],
    },
    "security": {
        "titles": ["SOC Analyst", "Security Analyst", "Cloud Security Engineer", "Senior Security Architect"],
        "core": ["Microsoft Sentinel", "SIEM", "Incident Response", "Azure Security", "IAM", "Threat Modeling"],
        "extra": ["Defender for Cloud", "Entra ID", "Zero Trust", "Splunk", "Python", "KQL",
                  "Vulnerability Management", "ISO 27001", "NIST CSF", "Penetration Testing"],
        "certs": ["Microsoft Certified: Security Operations Analyst Associate", "CISSP",
                  "CompTIA Security+", "Microsoft Certified: Cybersecurity Architect Expert"],
        "bullets": [
            ("Triaged {n}+ alerts per week in {a}, cutting mean time to respond by {p}%", ["Microsoft Sentinel", "Splunk", "SIEM"]),
            ("Rolled out {a} policies across {n} subscriptions", ["Defender for Cloud", "Entra ID", "Zero Trust"]),
            ("Wrote {a} detection rules that caught {n} real incidents in year one", ["KQL", "Microsoft Sentinel"]),
            ("Led {a} audit readiness, closing {p}% of findings ahead of schedule", ["ISO 27001", "NIST CSF"]),
        ],
    },
}

DEGREES = [
    ("BSc Computer Science", ["University of Toronto", "University of Waterloo", "McGill University",
                              "University of British Columbia", "York University"]),
    ("BEng Software Engineering", ["McMaster University", "Toronto Metropolitan University",
                                   "University of Ottawa", "Concordia University"]),
    ("BEng Electronics & Communication", ["Anna University", "VIT University", "Mumbai University"]),
    ("Diploma, Computer Programming", ["Sheridan College", "Seneca Polytechnic", "Humber Polytechnic",
                                       "George Brown College"]),
    ("BCom Information Systems", ["Western University", "Queen's University"]),
    ("Full-Stack / Data Bootcamp Certificate", ["Lighthouse Labs", "BrainStation", "Juno College"]),
]

SUMMARY_TAILS = [
    "Known for shipping reliable solutions and working closely with business stakeholders.",
    "Enjoys mentoring teammates and improving engineering practices.",
    "Focused on automation, clean design and measurable business impact.",
    "Comfortable in fast-paced agile teams and client-facing environments.",
    "Passionate about cloud-native architecture and continuous learning.",
]

# Job descriptions: what the agent will match candidates against.
JOBS = [
    {"id": "JD-001", "title": "Senior Data Engineer", "family": "data_engineer", "min_years": 5,
     "required": ["Python", "SQL", "Apache Spark", "Azure Data Factory"],
     "nice": ["Databricks", "Kafka", "dbt", "Terraform"]},
    {"id": "JD-002", "title": "AI Engineer (Generative AI)", "family": "ml_engineer", "min_years": 3,
     "required": ["Python", "LLMs", "Azure OpenAI", "RAG"],
     "nice": ["Azure AI Search", "LangChain", "Microsoft Foundry", "FastAPI", "Docker"]},
    {"id": "JD-003", "title": "DevOps Engineer", "family": "cloud_devops", "min_years": 3,
     "required": ["Azure", "Terraform", "Kubernetes", "CI/CD"],
     "nice": ["AKS", "GitHub Actions", "Helm", "Prometheus"]},
    {"id": "JD-004", "title": "Intermediate Java Developer", "family": "backend", "min_years": 2,
     "required": ["Java", "Spring Boot", "REST APIs", "SQL"],
     "nice": ["Microservices", "Kafka", "Azure", "Docker"]},
    {"id": "JD-005", "title": "Cloud Security Engineer", "family": "security", "min_years": 4,
     "required": ["Azure Security", "Microsoft Sentinel", "IAM", "Incident Response"],
     "nice": ["Defender for Cloud", "Entra ID", "KQL", "Zero Trust"]},
]

# Public-sector style Requests for Resumes (RFR): mandatory criteria are pass/fail, rated criteria
# earn points. The client is fictional. Used by the Week 4 submission workflow.
REQUESTS = [
    {"id": "RFR-001", "title": "Senior Data Engineer", "family": "data_engineer", "client": "Fictional Provincial Ministry (synthetic)", "term": "6 months, with option to extend 6 months", "location": "Toronto, ON (hybrid, 3 days on site)", "positions": 1, "mandatory": [{"id": "M1", "text": "Minimum 3 years of demonstrated experience building data pipelines with Apache Spark, Databricks or Azure Data Factory", "skills": ["Apache Spark", "Databricks", "Azure Data Factory"], "min_years": 3}, {"id": "M2", "text": "Demonstrated experience writing SQL or Python for data transformation", "skills": ["SQL", "Python"], "min_years": 0}], "rated": [{"id": "R1", "text": "Experience designing lakehouse solutions on Databricks, Delta Lake or Microsoft Fabric (4 years for full points)", "skills": ["Databricks", "Delta Lake", "Microsoft Fabric"], "target_years": 4, "points": 40}, {"id": "R2", "text": "Experience orchestrating pipelines with Azure Data Factory, Airflow or Microsoft Fabric (3 years for full points)", "skills": ["Azure Data Factory", "Airflow", "Microsoft Fabric"], "target_years": 3, "points": 35}, {"id": "R3", "text": "Experience with streaming data using Kafka (2 years for full points)", "skills": ["Kafka"], "target_years": 2, "points": 25}]},
    {"id": "RFR-002", "title": "Generative AI Developer", "family": "ml_engineer", "client": "Fictional Provincial Ministry (synthetic)", "term": "12 months", "location": "Toronto, ON (hybrid, 2 days on site)", "positions": 1, "mandatory": [{"id": "M1", "text": "Minimum 2 years of demonstrated experience building generative AI or retrieval solutions with Azure OpenAI, Microsoft Foundry, Azure AI Search or LangChain", "skills": ["Azure OpenAI", "Microsoft Foundry", "Azure AI Search", "LangChain"], "min_years": 2}, {"id": "M2", "text": "Demonstrated experience deploying models or APIs in containers (Docker or Kubernetes)", "skills": ["Docker", "Kubernetes"], "min_years": 0}], "rated": [{"id": "R1", "text": "Experience building retrieval-augmented generation (RAG) solutions with Azure AI Search (3 years for full points)", "skills": ["Azure AI Search"], "target_years": 3, "points": 40}, {"id": "R2", "text": "Experience with MLOps pipelines using Azure Machine Learning or MLflow (3 years for full points)", "skills": ["Azure Machine Learning", "MLflow", "MLOps"], "target_years": 3, "points": 35}, {"id": "R3", "text": "Experience serving models with FastAPI (2 years for full points)", "skills": ["FastAPI"], "target_years": 2, "points": 25}]},
    {"id": "RFR-003", "title": "Cloud Infrastructure Specialist", "family": "cloud_devops", "client": "Fictional Provincial Ministry (synthetic)", "term": "9 months, with option to extend 3 months", "location": "Toronto, ON (hybrid, 3 days on site)", "positions": 2, "mandatory": [{"id": "M1", "text": "Minimum 3 years of demonstrated experience with infrastructure as code using Terraform, Bicep or Ansible", "skills": ["Terraform", "Bicep", "Ansible"], "min_years": 3}, {"id": "M2", "text": "Demonstrated experience building CI/CD pipelines (GitHub Actions or Azure DevOps)", "skills": ["GitHub Actions", "Azure DevOps", "CI/CD"], "min_years": 0}], "rated": [{"id": "R1", "text": "Experience operating Kubernetes or AKS clusters in production (4 years for full points)", "skills": ["Kubernetes", "AKS"], "target_years": 4, "points": 50}, {"id": "R2", "text": "Experience implementing monitoring with Prometheus or Grafana (2 years for full points)", "skills": ["Prometheus", "Grafana"], "target_years": 2, "points": 50}]},
    {"id": "RFR-004", "title": "Security Operations Analyst", "family": "security", "client": "Fictional Provincial Ministry (synthetic)", "term": "6 months", "location": "Remote within Ontario", "positions": 1, "mandatory": [{"id": "M1", "text": "Minimum 2 years of demonstrated experience triaging alerts in Microsoft Sentinel, Splunk or another SIEM", "skills": ["Microsoft Sentinel", "Splunk", "SIEM"], "min_years": 2}], "rated": [{"id": "R1", "text": "Experience writing detection rules in KQL (3 years for full points)", "skills": ["KQL"], "target_years": 3, "points": 40}, {"id": "R2", "text": "Experience deploying Defender for Cloud, Entra ID or Zero Trust policies (3 years for full points)", "skills": ["Defender for Cloud", "Entra ID", "Zero Trust"], "target_years": 3, "points": 35}, {"id": "R3", "text": "Experience leading ISO 27001 or NIST CSF audit readiness (2 years for full points)", "skills": ["ISO 27001", "NIST CSF"], "target_years": 2, "points": 25}]},
]

# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------
@dataclass
class Job:
    title: str
    company: str
    start: str
    end: str
    bullets: list[str]


@dataclass
class Candidate:
    id: str
    name: str
    email: str
    phone: str
    city: str
    family: str
    years_experience: int
    headline: str
    summary: str
    skills: list[str]
    certifications: list[str]
    education: list[dict]
    experience: list[Job]
    # Fields a recruiter must NOT screen on. Present in some resumes on purpose
    # so we can test that the blind-screening step strips them.
    protected_info: dict = field(default_factory=dict)
    file: str = ""
    format: str = ""


# ---------------------------------------------------------------------------
# Generation
# ---------------------------------------------------------------------------
def month_str(d: date) -> str:
    return d.strftime("%b %Y")


def make_experience(rng: random.Random, fake: Faker, fam: dict, years: int, skills: list[str]) -> list[Job]:
    jobs, end = [], TODAY
    remaining_months = years * 12
    level = min(len(fam["titles"]) - 1, years // 4)
    while remaining_months > 0:
        length = min(remaining_months, rng.randint(14, 48))
        start_total = end.year * 12 + end.month - 1 - length
        start = date(start_total // 12, start_total % 12 + 1, 1)
        bullets = []
        for tmpl, options in rng.sample(fam["bullets"], k=rng.randint(2, 3)):
            # prefer a tool the candidate actually lists, so bullets stay consistent with the skills section
            owned = [o for o in options if o in skills]
            tool = rng.choice(owned) if owned else (options[0] if options else "")
            bullets.append(tmpl.format(a=tool, n=rng.randint(3, 50), p=rng.randint(15, 60)))
        jobs.append(Job(
            title=fam["titles"][max(level, 0)],
            company=fake.company(),
            start=month_str(start),
            end="Present" if end == TODAY else month_str(end),
            bullets=bullets,
        ))
        level -= 1
        remaining_months -= length
        # occasional career gap (realistic, and tests that gaps aren't over-penalized)
        gap = rng.choice([0, 0, 0, 0, 6])
        end_total = start_total - gap
        end = date(end_total // 12, end_total % 12 + 1, 1)
    return jobs


def make_candidate(i: int, rng: random.Random) -> Candidate:
    fake = Faker(rng.choice(LOCALES))
    fake.seed_instance(rng.randint(0, 10**9))
    family = rng.choice(list(FAMILIES))
    fam = FAMILIES[family]
    years = rng.choice([1, 2, 3, 3, 4, 5, 6, 7, 8, 10, 12, 15])

    n_core = rng.randint(3, len(fam["core"]))
    skills = rng.sample(fam["core"], n_core) + rng.sample(fam["extra"], rng.randint(2, 6))
    # a bit of cross-over so matching isn't trivially "same family = match"
    other = FAMILIES[rng.choice(list(FAMILIES))]
    skills += rng.sample(other["core"], 1)
    skills = list(dict.fromkeys(skills))

    certs = rng.sample(fam["certs"], k=rng.choice([0, 0, 1, 1, 2]) if len(fam["certs"]) > 1 else rng.choice([0, 1]))
    if years < 3:  # expert-level certs on a 1-year resume look fake; keep juniors cert-light
        certs = certs[:1] if "Expert" not in "".join(certs) and "CISSP" not in certs else []
    degree, schools = rng.choice(DEGREES)
    grad_year = TODAY.year - years - rng.randint(0, 2)
    education = [{"degree": degree, "school": rng.choice(schools), "year": grad_year}]

    name = f"{fake.first_name()} {fake.last_name()}"
    first = name.split()[0].lower()
    title = fam["titles"][min(len(fam["titles"]) - 1, years // 4)]

    protected = {}
    if rng.random() < 0.3:
        protected["date_of_birth"] = fake.date_of_birth(minimum_age=22 + years, maximum_age=30 + years).isoformat()
    if rng.random() < 0.15:
        protected["marital_status"] = rng.choice(["Married", "Single"])
    if rng.random() < 0.15:
        protected["photo_note"] = "[Photo]"

    return Candidate(
        id=f"CAND-{i:03d}",
        name=name,
        email=f"{first}.{rng.randint(10, 99)}@example.com",
        phone=f"+1 ({rng.randint(200, 999)}) 555-{rng.randint(1000, 9999)}",
        city=rng.choice(CITIES),
        family=family,
        years_experience=years,
        headline=title,
        summary=(f"{title} with {years}+ years of experience in {', '.join(skills[:3])}. "
                 + rng.choice(SUMMARY_TAILS)),
        skills=skills,
        certifications=certs,
        education=education,
        experience=make_experience(rng, fake, fam, years, skills),
        protected_info=protected,
    )


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------
def resume_lines(c: Candidate) -> list[tuple[str, str]]:
    """(style, text) pairs shared by the DOCX and PNG renderers."""
    lines = [("name", c.name), ("contact", f"{c.city} | {c.email} | {c.phone}")]
    if c.protected_info.get("photo_note"):
        lines.append(("contact", c.protected_info["photo_note"]))
    personal = [f"{k.replace('_', ' ').capitalize()}: {v}" for k, v in c.protected_info.items() if k != "photo_note"]
    if personal:
        lines.append(("contact", " | ".join(personal)))
    lines += [("h", "SUMMARY"), ("p", c.summary), ("h", "SKILLS"), ("p", ", ".join(c.skills))]
    lines.append(("h", "EXPERIENCE"))
    for j in c.experience:
        lines.append(("job", f"{j.title} - {j.company} ({j.start} - {j.end})"))
        lines += [("bullet", b) for b in j.bullets]
    lines.append(("h", "EDUCATION"))
    lines += [("p", f"{e['degree']}, {e['school']} ({e['year']})") for e in c.education]
    if c.certifications:
        lines.append(("h", "CERTIFICATIONS"))
        lines += [("bullet", x) for x in c.certifications]
    return lines


def render_pdf(c: Candidate, path: Path) -> None:
    ss = getSampleStyleSheet()
    styles = {
        "name": ParagraphStyle("n", parent=ss["Title"], fontSize=18, spaceAfter=2),
        "contact": ParagraphStyle("c", parent=ss["Normal"], fontSize=9, alignment=1),
        "h": ParagraphStyle("h", parent=ss["Heading3"], spaceBefore=8, spaceAfter=2),
        "p": ParagraphStyle("p", parent=ss["Normal"], fontSize=10, leading=13),
        "job": ParagraphStyle("j", parent=ss["Normal"], fontSize=10, fontName="Helvetica-Bold", spaceBefore=4),
    }
    story, bullets = [], []

    def flush():
        if bullets:
            story.append(ListFlowable([ListItem(Paragraph(b, styles["p"])) for b in bullets],
                                      bulletType="bullet", leftIndent=12))
            bullets.clear()

    for style, text in resume_lines(c):
        if style == "bullet":
            bullets.append(text)
            continue
        flush()
        story.append(Paragraph(text, styles[style]))
    flush()
    story.append(Spacer(1, 6))
    SimpleDocTemplate(str(path), pagesize=LETTER, topMargin=40, bottomMargin=40,
                      title=f"Resume - {c.name}").build(story)


def render_docx(c: Candidate, path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    for style, text in resume_lines(c):
        if style == "name":
            doc.add_heading(text, level=0)
        elif style == "h":
            doc.add_heading(text.title(), level=2)
        elif style == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        elif style == "job":
            doc.add_paragraph().add_run(text).bold = True
        else:
            doc.add_paragraph(text)
    doc.save(path)


def _font(size: int):
    for f in ["DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "Arial.ttf"]:
        try:
            return ImageFont.truetype(f, size)
        except OSError:
            continue
    return ImageFont.load_default()


def render_png(c: Candidate, path: Path, rng: random.Random) -> None:
    """Scanned-looking resume image: tests OCR / vision in Content Understanding."""
    import textwrap

    W, margin = 1275, 70  # ~150 dpi letter width
    fonts = {"name": _font(38), "h": _font(24), "job": _font(20), "default": _font(18)}
    rows = []
    for style, text in resume_lines(c):
        font = fonts.get(style, fonts["default"])
        prefix = "• " if style == "bullet" else ""
        for k, chunk in enumerate(textwrap.wrap(text, width=95 if style != "name" else 50)):
            rows.append((font, (prefix if k == 0 else "   ") + chunk, style))
    H = max(1650, margin * 2 + sum(f.size + 12 for f, _, _ in rows) + 40)
    img = Image.new("L", (W, H), 250)
    d = ImageDraw.Draw(img)
    y = margin
    for font, text, style in rows:
        if style == "h":
            y += 10
        d.text((margin, y), text, fill=25, font=font)
        y += font.size + 12
    # scanner artifacts: slight rotation, blur, noise
    img = img.rotate(rng.uniform(-1.2, 1.2), expand=False, fillcolor=245)
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    noise = Image.effect_noise(img.size, 12)
    img = Image.blend(img, noise, 0.06)
    img.convert("RGB").save(path, quality=85)


def render_job(j: dict, path: Path) -> None:
    fam = FAMILIES[j["family"]]
    text = f"""# {j['title']} ({j['id']})

**Location:** Toronto, ON (Hybrid, 3 days in office)
**Employment type:** Full-time, permanent

## About the role
We are looking for {'an' if j['title'][0] in 'AEIOU' else 'a'} {j['title']} to join a growing team building cloud-native products on Microsoft Azure.

## Responsibilities
""" + "\n".join(f"- Design, build and operate production solutions using {s}" for s in j["required"][:3]) + f"""
- Collaborate with product, data and security teams in an agile environment
- Document designs and share knowledge through code reviews and mentoring

## Requirements
- {j['min_years']}+ years of relevant professional experience
""" + "\n".join(f"- Hands-on experience with {s}" for s in j["required"]) + """

## Nice to have
""" + "\n".join(f"- {s}" for s in j["nice"]) + """

## Our commitment
We evaluate candidates on skills and experience only. We welcome applicants of all backgrounds.
"""
    path.write_text(text, encoding="utf-8")


def render_request(r: dict, path: Path) -> None:
    """A Request for Resumes as a PDF, laid out the way public-sector RFRs usually are."""
    ss = getSampleStyleSheet()
    p = ParagraphStyle("p", parent=ss["Normal"], fontSize=10, leading=13)
    story = [
        Paragraph(f"Request for Resumes {r['id']}: {r['title']}", ss["Title"]),
        Paragraph(f"<b>Client:</b> {r['client']}", p),
        Paragraph(f"<b>Number of resources:</b> {r['positions']}", p),
        Paragraph(f"<b>Term:</b> {r['term']}", p),
        Paragraph(f"<b>Work location:</b> {r['location']}", p),
        Spacer(1, 8),
        Paragraph("1. Mandatory requirements (pass/fail)", ss["Heading3"]),
        Paragraph("Resumes that do not demonstrate every mandatory requirement will not be evaluated further. "
                  "Experience must be shown in the resume with project or employer dates.", p),
    ]
    for c in r["mandatory"]:
        story.append(Paragraph(f"<b>{c['id']}.</b> {c['text']}", p))
    story += [Spacer(1, 8), Paragraph("2. Rated requirements", ss["Heading3"]),
              Paragraph("Points are awarded in proportion to demonstrated years of experience, "
                        "up to the stated number of years.", p)]
    for c in r["rated"]:
        story.append(Paragraph(f"<b>{c['id']}.</b> {c['text']} <b>({c['points']} points)</b>", p))
    total = sum(c["points"] for c in r["rated"])
    story += [Spacer(1, 6), Paragraph(f"<b>Total available points: {total}</b>", p), Spacer(1, 8),
              Paragraph("3. Submission instructions", ss["Heading3"]),
              Paragraph("For each requirement, state where in the resume the experience is demonstrated "
                        "(role, employer and dates). Evaluation is based on skills and experience only.", p)]
    SimpleDocTemplate(str(path), pagesize=LETTER, topMargin=40, bottomMargin=40,
                      title=f"{r['id']} {r['title']}").build(story)


def work_history(c: Candidate) -> list[dict]:
    """Jobs in the same shape the resume analyzer returns (JobTitle, Company, dates, Highlights)."""
    return [{"JobTitle": j.title, "Company": j.company, "StartDate": j.start, "EndDate": j.end,
             "Highlights": j.bullets} for j in c.experience]


def expected_screening(r: dict, candidates: list[Candidate]) -> list[dict]:
    """Ground truth for the submission workflow: apply the screening rules to the TRUE resume content.
    Evaluating the pipeline against this measures whether extraction kept what screening needs."""
    import sys
    sys.path.insert(0, str(ROOT / "src"))
    from copilot.screening import screen

    profiles = [{"id": c.id, "skills": c.skills, "work_history": work_history(c)} for c in candidates]
    return [{"id": x["candidate_id"], "compliance": x["compliance"], "points": x["points"]}
            for x in screen(r, profiles, today=TODAY)]


# ---------------------------------------------------------------------------
# Reference ranking (a simple, transparent baseline the agent can be evaluated against)
# ---------------------------------------------------------------------------
def baseline_score(c: Candidate, j: dict) -> float:
    skills = {s.lower() for s in c.skills}
    req = sum(s.lower() in skills for s in j["required"]) / len(j["required"])
    nice = sum(s.lower() in skills for s in j["nice"]) / len(j["nice"])
    exp = min(c.years_experience / j["min_years"], 1.5) / 1.5
    return round(0.6 * req + 0.2 * nice + 0.2 * exp, 3)


# ---------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--count", type=int, default=40)
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    rng = random.Random(args.seed)
    Faker.seed(args.seed)
    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "resumes").mkdir(parents=True)
    (OUT / "job-descriptions").mkdir(parents=True)

    candidates = []
    for i in range(1, args.count + 1):
        c = make_candidate(i, rng)
        fmt = rng.choices(["pdf", "docx", "png"], weights=[6, 3, 1])[0]
        c.format, c.file = fmt, f"resumes/{c.id}.{fmt}"
        path = OUT / c.file
        {"pdf": lambda: render_pdf(c, path),
         "docx": lambda: render_docx(c, path),
         "png": lambda: render_png(c, path, rng)}[fmt]()
        candidates.append(c)

    for j in JOBS:
        render_job(j, OUT / "job-descriptions" / f"{j['id']}.md")
    (OUT / "requests").mkdir()
    for r in REQUESTS:
        render_request(r, OUT / "requests" / f"{r['id']}.pdf")

    ground_truth = {
        "generated_with": {"seed": args.seed, "count": args.count},
        "candidates": [
            {k: v for k, v in asdict(c).items() if k not in ("summary", "experience")}
            | {"num_roles": len(c.experience), "experience": work_history(c)}
            for c in candidates
        ],
        "jobs": [
            j | {"baseline_top10": [
                {"id": c.id, "score": s} for s, c in
                sorted(((baseline_score(c, j), c) for c in candidates), key=lambda t: -t[0])[:10]
            ]}
            for j in JOBS
        ],
        # expected screening outcome per request, from the true resume content (see copilot.screening)
        "requests": [r | {"expected": expected_screening(r, candidates)} for r in REQUESTS],
    }
    (OUT / "ground_truth.json").write_text(json.dumps(ground_truth, indent=2), encoding="utf-8")

    fmts = {f: sum(c.format == f for c in candidates) for f in ("pdf", "docx", "png")}
    prot = sum(bool(c.protected_info) for c in candidates)
    print(f"Generated {len(candidates)} resumes {fmts}, {len(JOBS)} job descriptions, "
          f"{len(REQUESTS)} requests for resumes -> {OUT}")
    print(f"{prot} resumes include protected info (DOB / marital status / photo) for blind-screening tests")


if __name__ == "__main__":
    main()

"""
Submission drafts for a Request for Resumes: a response grid per candidate, a short summary,
and a human approval step before anything leaves the agency.

Rules this module enforces:
- The grid is built from evidence in code; every response points to a role, employer and dates.
- The model may write the summary, but only from the grid's facts. A grounding check rejects any
  summary that adds numbers or skills the evidence doesn't contain, and falls back to a template.
- Drafts are blind (candidate ID only). The name is added back only when an approved draft is
  exported, so screening can never be influenced by it.
"""
from __future__ import annotations

import json
import math
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from copilot.screening import MET, NOT_MET, VERIFY

ROOT = Path(__file__).resolve().parents[2]
SUBMISSIONS = ROOT / "data" / "submissions"
MAX_JOBS_PER_ROW = 3
PROTECTED_TERMS = re.compile(r"\b(age|aged|born|birth|married|single|marital|gender|nationality|religion|photo)\b", re.I)


class ApprovalRequired(RuntimeError):
    pass


# ---------------------------------------------------------------------------
# Response grid
# ---------------------------------------------------------------------------
def _years(months: int) -> str:
    y = round(months / 12, 1)
    return f"{y:g} year{'s' if y != 1 else ''}"


def criterion_response(row: dict[str, Any]) -> str:
    """Plain-language response for one criterion, citing where the experience is shown."""
    if row["status"] == VERIFY:
        return (f"{', '.join(row['listed_in_skills'])} is listed in the candidate's skills but not shown in a dated "
                "role. Recruiter to confirm with the candidate before submitting.")
    if not row["evidence"]:
        return "Not demonstrated in the resume."
    jobs = row["evidence"][:MAX_JOBS_PER_ROW]
    cites = "; ".join(f"{j['title']}, {j['company']} ({j['start']} - {j['end']}): \"{j['highlights'][0]}\""
                      for j in jobs)
    more = f" (+{len(row['evidence']) - len(jobs)} more)" if len(row["evidence"]) > len(jobs) else ""
    lead = f"Demonstrated in {len(row['evidence'])} role{'s' if len(row['evidence']) > 1 else ''}, {_years(row['evidence_months'])} in total"
    if row["kind"] == "mandatory" and row["status"] == NOT_MET and row.get("required_years"):
        lead += f", below the {row['required_years']} years required"
    return f"{lead}. {cites}{more}."


def build_draft(requirement: dict, evaluation: dict, profile: dict) -> dict[str, Any]:
    rows, warnings = [], []
    for r in evaluation["mandatory"] + evaluation["rated"]:
        row = {
            "id": r["id"], "kind": r["kind"], "text": r["text"], "status": r["status"],
            "evidence_years": r["evidence_years"], "response": criterion_response(r),
            "evidence": r["evidence"],
        }
        if r["kind"] == "rated":
            row |= {"points_awarded": r["points_awarded"], "points_max": r["points_max"]}
        rows.append(row)
        if r["status"] == VERIFY:
            warnings.append(f"{r['id']}: {', '.join(r['listed_in_skills'])} listed but not shown in a dated role. "
                            "Confirm with the candidate and add the project to the resume.")
        elif r["kind"] == "mandatory" and r["status"] == NOT_MET:
            warnings.append(f"{r['id']}: mandatory requirement not demonstrated. This resume would be screened out.")
    return {
        "requirement_id": requirement["id"],
        "requirement_title": requirement.get("title"),
        "candidate_id": profile["id"],
        "candidate_headline": profile.get("headline"),
        "candidate_years": profile.get("years_experience"),
        "compliance": evaluation["compliance"],
        "points": evaluation["points"],
        "points_max": evaluation["points_max"],
        "rows": rows,
        "warnings": warnings,
        "summary": template_summary(evaluation, profile),
        "summary_source": "template",
        "status": "draft",
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


def template_summary(evaluation: dict, profile: dict) -> str:
    mand = evaluation["mandatory"]
    met = sum(m["status"] == MET for m in mand)
    verify = [m["id"] for m in mand if m["status"] == VERIFY]
    rated = sorted(evaluation["rated"], key=lambda r: -r["points_awarded"])
    strong = [r["id"] for r in rated if r["points_awarded"] >= 0.75 * r["points_max"] and r["points_max"]]
    gaps = [r["id"] for r in rated if r["points_awarded"] == 0]
    mand_text = f"Meets {met} of {len(mand)} mandatory requirements"
    if verify:
        mand_text += f" ({', '.join(verify)} needs verification)"
    parts = [f"{profile.get('headline') or 'Candidate'} with {profile.get('years_experience')} years of experience.",
             f"{mand_text}; scores {evaluation['points']:g} of {evaluation['points_max']:g} rated points."]
    if strong:
        parts.append(f"Strongest on {', '.join(strong)}.")
    if gaps:
        parts.append(f"No demonstrated experience for {', '.join(gaps)}.")
    return " ".join(parts)


# ---------------------------------------------------------------------------
# Model-written summary, with a grounding check
# ---------------------------------------------------------------------------
SUMMARY_INSTRUCTIONS = """\
You write the 2-3 sentence candidate summary that goes at the top of a staffing submission.
Use ONLY the facts provided. Do not add skills, employers, numbers or achievements that are not in
the facts. Refer to the person as "the candidate". Never mention age, gender, family status,
nationality or anything not job-related. If a mandatory requirement is not met or needs
verification, say so plainly. Return only the summary text."""


def _facts(draft: dict) -> dict[str, Any]:
    return {
        "role": draft["requirement_title"], "current_title": draft["candidate_headline"],
        "years_experience": draft["candidate_years"], "compliance": draft["compliance"],
        "rated_points": f"{draft['points']:g} of {draft['points_max']:g}",
        "criteria": [{k: r.get(k) for k in ("id", "kind", "text", "status", "evidence_years",
                                            "points_awarded", "points_max", "response")} for r in draft["rows"]],
    }


def check_grounded(summary: str, facts: dict, vocabulary: list[str]) -> list[str]:
    """Problems with a model-written summary; empty list means it's safe to use."""
    problems = []
    source = json.dumps(facts)
    values = [float(n) for n in re.findall(r"\d+(?:\.\d+)?", source)]
    # "more than 5 years" is a fair reading of 5.2 years, so floors and roundings are allowed too
    allowed = {f"{x:g}" for v in values for x in (v, math.floor(v), round(v))}
    for n in re.findall(r"\d+(?:\.\d+)?", summary):
        if f"{float(n):g}" not in allowed:
            problems.append(f"number not in evidence: {n}")
    for skill in vocabulary:
        if re.search(rf"(?<![\w]){re.escape(skill)}(?![\w])", summary, re.I) and skill.lower() not in source.lower():
            problems.append(f"skill not in evidence: {skill}")
    if PROTECTED_TERMS.search(summary):
        problems.append("mentions a protected or non-job-related attribute")
    if not summary.strip():
        problems.append("empty")
    return problems


async def write_summary(draft: dict, client, vocabulary: list[str]) -> dict:
    """Ask the model for a summary; keep it only if it passes the grounding check."""
    from agent_framework import Agent

    facts = _facts(draft)
    agent = Agent(client=client, name="submission-writer", instructions=SUMMARY_INSTRUCTIONS)
    reply = await agent.run("Facts:\n" + json.dumps(facts, indent=2))
    text = (reply.text or "").strip()
    problems = check_grounded(text, facts, vocabulary)
    if problems:
        draft["summary_rejected"] = {"text": text, "problems": problems}
    else:
        draft["summary"], draft["summary_source"] = text, "model"
    return draft


# ---------------------------------------------------------------------------
# Review workflow: save -> recruiter approves or rejects -> export
# ---------------------------------------------------------------------------
def draft_path(requirement_id: str, candidate_id: str, folder: Path | None = None) -> Path:
    return (folder or SUBMISSIONS) / requirement_id / f"{candidate_id}.json"


def to_markdown(draft: dict) -> str:
    lines = [f"# DRAFT submission: {draft['requirement_id']} {draft['requirement_title']}", "",
             f"Candidate: **{draft['candidate_id']}** (name added on export after approval)  ",
             f"Status: **{draft['status']}** · Mandatory: **{draft['compliance']}** · "
             f"Rated: **{draft['points']:g} / {draft['points_max']:g}**", "",
             f"> {draft['summary']}", ""]
    if draft["warnings"]:
        lines += ["## Check before submitting", ""] + [f"- {w}" for w in draft["warnings"]] + [""]
    lines += ["## Response grid", "", "| Req | Result | Response |", "|---|---|---|"]
    for r in draft["rows"]:
        result = r["status"] if r["kind"] == "mandatory" else f"{r['points_awarded']:g}/{r['points_max']:g}"
        lines.append(f"| {r['id']} | {result} | {r['response'].replace('|', '/')} |")
    return "\n".join(lines) + "\n"


def save_draft(draft: dict, folder: Path | None = None) -> Path:
    path = draft_path(draft["requirement_id"], draft["candidate_id"], folder)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(draft, indent=2), encoding="utf-8")
    path.with_suffix(".md").write_text(to_markdown(draft), encoding="utf-8")
    return path


def review(path: Path, reviewer: str, decision: str, note: str = "") -> dict:
    """Record a recruiter's decision ('approved' or 'rejected') on a saved draft."""
    if decision not in ("approved", "rejected"):
        raise ValueError("decision must be 'approved' or 'rejected'")
    if not reviewer.strip():
        raise ValueError("a reviewer name is required")
    draft = json.loads(path.read_text(encoding="utf-8"))
    draft |= {"status": decision, "reviewer": reviewer.strip(), "review_note": note,
              "reviewed_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}
    path.write_text(json.dumps(draft, indent=2), encoding="utf-8")
    path.with_suffix(".md").write_text(to_markdown(draft), encoding="utf-8")
    return draft


def export_docx(draft: dict, profile: dict, out: Path) -> Path:
    """Client-ready Word document. Refuses unless a recruiter approved the draft."""
    if draft.get("status") != "approved":
        raise ApprovalRequired(f"{draft['candidate_id']} for {draft['requirement_id']} is '{draft.get('status')}'. "
                               "A recruiter must approve it first (scripts/review_submission.py).")
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(f"Submission for {draft['requirement_id']}: {draft['requirement_title']}", level=1)
    doc.add_paragraph().add_run(f"Candidate: {profile.get('name') or draft['candidate_id']}").bold = True
    doc.add_paragraph(draft["summary"])
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Requirement", "Response", "Result")):
        cell.text = text
    for r in draft["rows"]:
        cells = table.add_row().cells
        cells[0].text = f"{r['id']}. {r['text']}"
        cells[1].text = r["response"]
        cells[2].text = (r["status"].replace("_", " ").title() if r["kind"] == "mandatory"
                         else f"{r['points_awarded']:g} / {r['points_max']:g}")
    doc.add_paragraph()
    doc.add_paragraph().add_run(
        f"Reviewed and approved by {draft['reviewer']} on {draft['reviewed_at'][:10]}.").italic = True
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out
